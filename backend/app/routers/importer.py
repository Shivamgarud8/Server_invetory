"""Import the legacy Excel sheets (or CSV) into servers, credentials and versions.

Headers are matched loosely ("Server Name", "server_name", "SERVER" all work). Every import can be
run as a dry run first; the real run is all-or-nothing.
"""

import csv
import io
import ipaddress
import re
from datetime import date, datetime, timezone

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..audit import audit
from ..db import get_db
from ..deps import Principal, require
from ..models import (
    Application,
    ApplicationVersion,
    Credential,
    CredentialType,
    Environment,
    Server,
    ServerNetwork,
    ServerSpec,
    Tag,
    VersionHistory,
)
from ..security import VaultError, get_vault

router = APIRouter(prefix="/api/import", tags=["import"])

MAX_BYTES = 10 * 1024 * 1024

# canonical field -> accepted header aliases (normalised: lowercase alnum only)
COLUMNS = {
    "servers": {
        "name": ["servername", "server", "name", "host"],
        "hostname": ["hostname", "fqdn"],
        "environment": ["environment", "env"],
        "server_type": ["servertype", "type", "role"],
        "private_ip": ["privateip", "ip", "ipaddress", "internalip"],
        "public_ip": ["publicip", "externalip"],
        "network_ip": ["networkip", "vpnip"],
        "provider": ["provider", "cloud"],
        "location": ["location", "region", "datacenter"],
        "os_name": ["os", "operatingsystem", "osname"],
        "os_version": ["osversion"],
        "kernel_version": ["kernel", "kernelversion"],
        "architecture": ["architecture", "arch"],
        "cpu_cores": ["cpu", "cpucores", "cores", "vcpu"],
        "ram_gb": ["ram", "ramgb", "memory"],
        "disk_gb": ["disk", "diskgb", "storage"],
        "disk_usage_pct": ["diskusage", "diskusagepct", "diskused"],
        "docker_version": ["docker", "dockerversion"],
        "compose_version": ["compose", "composeversion", "dockercomposeversion"],
        "status": ["status", "serverstatus"],
        "owner": ["owner"],
        "purpose": ["purpose", "services", "usage"],
        "notes": ["notes", "remarks", "comment", "comments"],
        "tags": ["tags", "tag", "client"],
    },
    "credentials": {
        "name": ["credentialname", "credential", "name", "title"],
        "type": ["type", "credentialtype", "servicetype"],
        "category": ["category"],
        "server": ["server", "servername"],
        "environment": ["environment", "env"],
        "host": ["ip", "ipaddress", "host", "hostip"],
        "port": ["port"],
        "username": ["username", "user", "login"],
        "password": ["password", "pass", "pwd", "secret"],
        "url": ["url", "link", "endpoint"],
        "service": ["service"],
        "description": ["description", "desc"],
        "owner": ["owner"],
        "status": ["status"],
        "expires_at": ["expirydate", "expiry", "expires", "expiresat"],
        "notes": ["notes", "remarks", "comment", "comments"],
        "tags": ["tags", "tag", "client"],
    },
    "versions": {
        "application": ["application", "app", "applicationname"],
        "component": ["component", "service", "module"],
        "server": ["server", "servername"],
        "environment": ["environment", "env"],
        "version": ["version", "currentversion"],
        "previous_version": ["previousversion", "oldversion"],
        "image": ["image", "dockerimage"],
        "deployed_at": ["deploymentdate", "deployed", "deployedat", "updated", "date"],
        "notes": ["notes", "remarks", "comment", "comments"],
    },
}

REQUIRED = {
    "servers": ["name", "environment"],
    "credentials": ["name", "environment"],
    "versions": ["application", "environment", "version"],
}

ENV_ALIASES = {"production": "PROD", "prd": "PROD", "stage": "STAGING", "stg": "STAGING",
               "development": "DEV", "qa": "TEST", "testing": "TEST"}

PERMS = {"servers": "servers.create", "credentials": "credentials.create", "versions": "versions.create"}


def _norm(h) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def _read_rows(content: bytes, filename: str, sheet: str | None) -> list[dict]:
    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        return [dict(r) for r in csv.DictReader(io.StringIO(text))]
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    headers = None
    rows = []
    for raw in it:
        if headers is None:
            if raw and sum(1 for c in raw if c not in (None, "")) >= 2:  # first non-empty row is the header
                headers = [str(c or "").strip() for c in raw]
            continue
        if raw and any(c not in (None, "") for c in raw):
            rows.append(dict(zip(headers, raw)))
    wb.close()
    return rows


def _map(kind: str, row: dict) -> dict:
    aliases = COLUMNS[kind]
    normed = {_norm(k): v for k, v in row.items()}
    out = {}
    for field, names in aliases.items():
        for n in names:
            if n in normed and normed[n] not in (None, ""):
                v = normed[n]
                out[field] = v.strip() if isinstance(v, str) else v
                break
    return out


def _num(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"\d+(?:\.\d+)?", str(v))
    return float(m.group()) if m else None


def _date(v) -> date | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unrecognised date '{v}'")


def _ips(v) -> list[str]:
    out = []
    for part in re.split(r"[,\s;/]+", str(v or "")):
        part = part.strip()
        if not part or part == "-":
            continue
        try:
            ipaddress.ip_address(part)
            out.append(part)
        except ValueError:
            raise ValueError(f"invalid IP '{part}'")
    return out


class Ctx:
    def __init__(self, db: Session, p: Principal):
        self.db = db
        self.p = p
        self.envs = {e.name.upper(): e for e in db.scalars(select(Environment))}
        self.tags = {t.name.upper(): t for t in db.scalars(select(Tag))}
        self.types = {t.name.lower(): t for t in db.scalars(select(CredentialType))}
        self.servers = {s.name.lower(): s for s in db.scalars(select(Server))}
        self.apps = {a.name.lower(): a for a in db.scalars(select(Application))}

    def env(self, v) -> Environment:
        key = str(v).strip().upper()
        key = ENV_ALIASES.get(key.lower(), key)
        env = self.envs.get(key)
        if not env:
            raise ValueError(f"unknown environment '{v}' (valid: {', '.join(self.envs)})")
        if not self.p.can_env(env.id):
            raise ValueError(f"no access to environment {env.name}")
        return env

    def tag_list(self, v) -> list[Tag]:
        out = []
        for name in re.split(r"[,\s;]+", str(v or "")):
            name = name.strip().upper()
            if not name:
                continue
            if name not in self.tags:
                t = Tag(name=name)
                self.db.add(t)
                self.tags[name] = t
            out.append(self.tags[name])
        return out

    def server(self, v) -> Server | None:
        if v in (None, "", "-"):
            return None
        s = self.servers.get(str(v).strip().lower())
        if not s:
            raise ValueError(f"unknown server '{v}' (import servers first)")
        return s


def _import_server(ctx: Ctx, r: dict) -> str:
    env = ctx.env(r["environment"])
    networks = (
        [("private", ip) for ip in _ips(r.get("private_ip"))]
        + [("public", ip) for ip in _ips(r.get("public_ip"))]
        + [("network", ip) for ip in _ips(r.get("network_ip"))]
    )
    name = str(r["name"]).strip()
    s = ctx.servers.get(name.lower())
    action = "update" if s else "create"
    if s is None:
        s = Server(name=name, created_by_id=ctx.p.user.id)
        ctx.db.add(s)
        ctx.servers[name.lower()] = s
    elif not ctx.p.can_env(s.environment_id):
        raise ValueError("no access to existing server's environment")
    s.environment = env
    s.updated_by_id = ctx.p.user.id
    for f in ("hostname", "provider", "location", "os_name", "os_version", "kernel_version", "architecture",
              "docker_version", "compose_version", "owner", "purpose", "notes"):
        if f in r:
            setattr(s, f, str(r[f]))
    s.server_type = str(r.get("server_type") or s.server_type or "Server")
    s.status = str(r.get("status") or s.status or "Active")
    if s.spec is None:
        s.spec = ServerSpec()
    cpu = _num(r.get("cpu_cores"))
    s.spec.cpu_cores = int(cpu) if cpu is not None else s.spec.cpu_cores
    for f in ("ram_gb", "disk_gb", "disk_usage_pct"):
        v = _num(r.get(f))
        if v is not None:
            setattr(s.spec, f, v)
    if networks:
        s.networks = [ServerNetwork(ip_type=t, ip_address=ip) for t, ip in networks]
    if "tags" in r:
        s.tags = ctx.tag_list(r["tags"])
    return action


def _import_credential(ctx: Ctx, r: dict) -> tuple[str, Credential]:
    env = ctx.env(r["environment"])
    server = ctx.server(r.get("server"))
    ctype = ctx.types.get(str(r.get("type") or "other").strip().lower()) or ctx.types.get("other")
    port = _num(r.get("port"))
    name = str(r["name"]).strip()
    username = str(r["username"]) if r.get("username") is not None else None
    existing = ctx.db.scalar(select(Credential).where(
        Credential.name == name, Credential.environment_id == env.id,
        Credential.username.is_(None) if username is None else Credential.username == username))
    c = existing or Credential(created_by_id=ctx.p.user.id, secret_provider=get_vault().name)
    if existing is None:
        ctx.db.add(c)
    c.name, c.environment, c.server, c.credential_type = name, env, server, ctype
    c.username = username
    c.port = int(port) if port else (ctype.default_port if ctype and not c.port else c.port)
    c.updated_by_id = ctx.p.user.id
    c.host = str(r["host"]) if r.get("host") else (c.host or (str(server.networks[0].ip_address) if server and server.networks else None))
    for f in ("category", "url", "service", "description", "owner", "notes"):
        if f in r:
            setattr(c, f, str(r[f]))
    if r.get("type") and ctype and ctype.name.lower() != str(r["type"]).strip().lower() and not c.category:
        c.category = str(r["type"])  # keep the original label when it fell back to "Other"
    c.status = str(r.get("status") or c.status or "Active")
    if "expires_at" in r:
        c.expires_at = _date(r["expires_at"])
    if "tags" in r:
        c.tags = ctx.tag_list(r["tags"])
    return ("update" if existing else "create"), c


def _import_version(ctx: Ctx, r: dict, who: str) -> str:
    env = ctx.env(r["environment"])
    server = ctx.server(r.get("server"))
    app_name = str(r["application"]).strip()
    app = ctx.apps.get(app_name.lower())
    if app is None:
        app = Application(name=app_name)
        ctx.db.add(app)
        ctx.apps[app_name.lower()] = app
    component = str(r.get("component") or app_name).strip()
    version = str(r["version"]).strip()
    deployed = _date(r.get("deployed_at"))
    deployed_at = datetime.combine(deployed, datetime.min.time(), tzinfo=timezone.utc) if deployed else datetime.now(timezone.utc)

    d = None
    if app.id:
        d = ctx.db.scalar(select(ApplicationVersion).where(
            ApplicationVersion.application_id == app.id, ApplicationVersion.component == component,
            ApplicationVersion.environment_id == env.id,
            ApplicationVersion.server_id.is_(None) if server is None else ApplicationVersion.server_id == server.id))
    if d and d.current_version == version:
        return "unchanged"
    old = d.current_version if d else (str(r["previous_version"]) if r.get("previous_version") else None)
    if d is None:
        d = ApplicationVersion(application=app, component=component, server=server, environment=env,
                               current_version=version)
        ctx.db.add(d)
        action = "create"
    else:
        action = "update"
    d.previous_version = old
    d.current_version = version
    d.deployed_at = deployed_at
    d.updated_by_id = ctx.p.user.id
    if r.get("image"):
        d.image = str(r["image"])
    if r.get("notes"):
        d.notes = str(r["notes"])
    ctx.db.flush()
    ctx.db.add(VersionHistory(
        application_version_id=d.id, application_name=app.name, component=component,
        server_name=server.name if server else None, environment_name=env.name, old_version=old,
        new_version=version, changed_by=who, changed_at=deployed_at, notes="Imported",
    ))
    return action


@router.post("/{kind}")
async def run_import(
    kind: str, request: Request, file: UploadFile = File(...), dry_run: bool = Form(True),
    sheet: str | None = Form(None), db: Session = Depends(get_db), p: Principal = Depends(require("import.run")),
):
    if kind not in COLUMNS:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Unknown import type")
    if not p.can(PERMS[kind]):
        raise HTTPException(status.HTTP_403_FORBIDDEN, f"Missing permission: {PERMS[kind]}")
    content = await file.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "File larger than 10 MB")
    try:
        raw_rows = _read_rows(content, file.filename or "", sheet)
    except Exception:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Could not read file — upload .xlsx or .csv")

    ctx = Ctx(db, p)
    results, errors, secrets = [], [], []
    counts = {"create": 0, "update": 0, "unchanged": 0}
    for i, raw in enumerate(raw_rows, start=2):
        r = _map(kind, raw)
        missing = [f for f in REQUIRED[kind] if not r.get(f)]
        if missing:
            errors.append({"row": i, "error": f"missing {', '.join(missing)}"})
            continue
        try:
            with db.begin_nested():
                if kind == "servers":
                    action = _import_server(ctx, r)
                elif kind == "credentials":
                    action, cred = _import_credential(ctx, r)
                else:
                    action = _import_version(ctx, r, p.user.username)
                db.flush()
        except Exception as e:  # noqa: BLE001 — report row-level problems back to the user
            msg = str(e).split("\n")[0][:200]
            errors.append({"row": i, "error": msg})
            # nested rollback may have expired cached objects
            ctx = Ctx(db, p)
            continue
        counts[action] += 1
        if kind == "credentials" and r.get("password"):
            secrets.append((cred, str(r["password"])))
        label = r.get("name") or f"{r.get('application')} / {r.get('component') or ''}"
        results.append({"row": i, "action": action, "label": str(label),
                        "has_password": bool(kind == "credentials" and r.get("password"))})

    unmapped = sorted({k for raw in raw_rows[:1] for k in raw if _map(kind, {k: "x"}) == {}})
    summary = {"rows": len(raw_rows), "counts": counts, "errors": errors, "results": results[:500],
               "unmapped_columns": unmapped, "dry_run": dry_run}

    if dry_run or errors:
        db.rollback()
        summary["dry_run"] = True
        if errors and not dry_run:
            summary["message"] = "Nothing imported — fix the errors and run again"
        return summary

    try:
        for c, secret in secrets:
            c.secret_enc = get_vault().store(c.id, secret)
            c.secret_provider = get_vault().name
            c.has_secret = True
            c.password_changed_at = datetime.now(timezone.utc)
    except VaultError:
        db.rollback()
        raise HTTPException(status.HTTP_502_BAD_GATEWAY, "Secret storage unavailable — nothing imported")
    audit(db, request, f"IMPORT_{kind.upper()}", p.user, entity_label=file.filename,
          details={"rows": len(raw_rows), **counts, "passwords": len(secrets)})
    db.commit()
    return summary


@router.get("/{kind}/template")
def template(kind: str, _: Principal = Depends(require("import.run"))):
    if kind not in COLUMNS:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Unknown import type")
    headers = {
        "servers": ["Server Name", "Hostname", "Environment", "Server Type", "Private IP", "Public IP", "Network IP",
                    "Provider", "Location", "OS", "OS Version", "Kernel", "Architecture", "CPU", "RAM", "Disk",
                    "Disk Usage", "Docker Version", "Compose Version", "Status", "Owner", "Purpose", "Tags", "Notes"],
        "credentials": ["Credential Name", "Type", "Category", "Server", "Environment", "IP", "Port", "Username",
                        "Password", "URL", "Service", "Description", "Owner", "Status", "Expiry Date", "Tags", "Notes"],
        "versions": ["Application", "Component", "Server", "Environment", "Version", "Previous Version", "Image",
                     "Deployment Date", "Notes"],
    }[kind]
    example = {
        "servers": ["srv-p-03-db", "srv-p-03-db.internal", "PROD", "Database Server", "10.80.0.13", "", "", "Hetzner",
                    "Mumbai", "Ubuntu", "24.04", "6.8.0", "x86_64", "4", "16", "300", "42", "29.0.1", "2.40.3",
                    "Active", "Shivam", "PostgreSQL / Redis / RabbitMQ / MinIO", "DATABASE CRITICAL", ""],
        "credentials": ["srv-p-03-db Root", "SSH", "Server", "srv-p-03-db", "PROD", "10.80.0.13", "22", "root",
                        "", "", "ssh", "Production database server root access", "Shivam", "Active", "", "CRITICAL", ""],
        "versions": ["TagID API", "Admin API", "srv-p-01-app", "PROD", "v3.33.4", "v3.33.3", "", "2026-10-08", ""],
    }[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = kind.capitalize()
    ws.append(headers)
    ws.append(example)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="tagid-{kind}-template.xlsx"'},
    )
